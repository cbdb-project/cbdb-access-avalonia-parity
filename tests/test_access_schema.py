"""Tests for cbdb_parity.access_schema — the TablesFields.xlsx loader."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from cbdb_parity.access_schema import (
    AccessSchemaError,
    AccessTable,
    load_access_schema,
)


def _write_sheet(path: Path, rows: list[tuple[object, ...]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "TablesFields"
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


_HEADER = (
    "RowNum", "DumpTblNm", "DumpFldNm", "AccessTblNm", "AccessFldNm",
    "IndexOnField", "DataFormat", "NULL_allowed", "ForeignKey", "ForeignKeyBaseField",
)


def test_load_minimal_two_tables(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "BIOG_MAIN", "c_personid", "Primary Key", "Long", False, None, None),
        (None, None, None, "BIOG_MAIN", "c_name_chn", None, "Text", True, None, None),
        (None, None, None, "ADDR_CODES", "c_addr_id", "Primary Key", "Long", False, None, None),
    ])
    schema = load_access_schema(path)
    assert len(schema) == 2
    assert "BIOG_MAIN" in schema.tables
    assert "ADDR_CODES" in schema.tables
    bm: AccessTable = schema["BIOG_MAIN"]
    assert [c.name for c in bm.columns] == ["c_personid", "c_name_chn"]
    assert bm.primary_key_columns == ("c_personid",)
    pk_col = bm.columns[0]
    assert pk_col.is_primary_key
    assert pk_col.data_format == "Long"
    assert pk_col.nullable is False


def test_load_normalises_primary_and_primary_key_labels(tmp_path: Path) -> None:
    """The real xlsx uses both 'Primary' and 'Primary Key' as PK markers."""
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", "Primary",     "Long", False, None, None),
        (None, None, None, "T", "b", "Primary Key", "Long", False, None, None),
        (None, None, None, "T", "c", None,          "Long", True,  None, None),
    ])
    schema = load_access_schema(path)
    assert schema["T"].primary_key_columns == ("a", "b")


def test_load_normalises_binary_typo(tmp_path: Path) -> None:
    """The real xlsx has both 'Binary' and 'binary'; both must normalise."""
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "Binary", True, None, None),
        (None, None, None, "T", "b", None, "binary", True, None, None),
    ])
    schema = load_access_schema(path)
    cols = schema["T"].columns
    assert cols[0].data_format == "Binary"
    assert cols[1].data_format == "Binary"


def test_load_records_foreign_key_target(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "ADDR_BELONGS_DATA", "c_addr_id", "Primary Key", "Long", False, "ADDR_CODES", "c_addr_id"),
    ])
    schema = load_access_schema(path)
    col = schema["ADDR_BELONGS_DATA"].columns[0]
    assert col.foreign_key_table == "ADDR_CODES"
    assert col.foreign_key_column == "c_addr_id"


def test_unknown_data_format_raises(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "QuantumString", True, None, None),
    ])
    with pytest.raises(AccessSchemaError, match=r"unknown DataFormat"):
        load_access_schema(path)


def test_duplicate_column_in_same_table_raises(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "Long", True, None, None),
        (None, None, None, "T", "a", None, "Text", True, None, None),
    ])
    with pytest.raises(AccessSchemaError, match=r"duplicate column"):
        load_access_schema(path)


def test_blank_rows_are_skipped(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "Long", True, None, None),
        (None, None, None, None, None, None, None, None, None, None),  # blank-ish
        (None, None, None, "T", "b", None, "Long", True, None, None),
    ])
    schema = load_access_schema(path)
    assert [c.name for c in schema["T"].columns] == ["a", "b"]


def test_missing_xlsx_raises(tmp_path: Path) -> None:
    with pytest.raises(AccessSchemaError, match=r"not found"):
        load_access_schema(tmp_path / "missing.xlsx")


def test_bad_header_raises(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        ("foo", "bar", "baz"),
    ])
    with pytest.raises(AccessSchemaError, match=r"unexpected header"):
        load_access_schema(path)


@pytest.mark.parametrize(
    "cell, expected",
    [
        (True, True),
        (False, False),
        (1, True),
        (0, False),
        ("Yes", True),
        ("No", False),
        ("yes", True),
        ("NO", False),
        ("true", True),
        ("false", False),
        ("", True),    # blank cell → default nullable
        (None, True),
    ],
)
def test_nullable_cell_interpretations(cell: object, expected: bool, tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "Long", cell, None, None),
    ])
    schema = load_access_schema(path)
    assert schema["T"].columns[0].nullable is expected


def test_nullable_garbage_raises(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "Long", "maybe", None, None),
    ])
    with pytest.raises(AccessSchemaError, match=r"cannot interpret nullable"):
        load_access_schema(path)


def test_fk_whitespace_treated_as_none(tmp_path: Path) -> None:
    """A foreign-key cell containing only whitespace must NOT be stored as
    an empty FK target — downstream code uses `is None` to detect 'no FK'."""
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "Long", True, "   ", "   "),
    ])
    schema = load_access_schema(path)
    col = schema["T"].columns[0]
    assert col.foreign_key_table is None
    assert col.foreign_key_column is None


def test_blank_data_format_stored_as_none(tmp_path: Path) -> None:
    """Real xlsx has ~20 blank-DataFormat rows (e.g.
    POSTED_TO_ADDR_DATA.c_personid). The loader keeps the column and
    stores `data_format=None` so the mdb writer (Phase 1.3b) can fall
    back to the Datadump's CREATE TABLE for that column's type."""
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "c_personid", None, None, True, None, None),
        (None, None, None, "T", "c_name", None, "Text", True, None, None),
    ])
    schema = load_access_schema(path)
    cols = schema["T"].columns
    assert cols[0].name == "c_personid"
    assert cols[0].data_format is None
    assert cols[1].name == "c_name"
    assert cols[1].data_format == "Text"


def test_extra_trailing_header_columns_raises(tmp_path: Path) -> None:
    """An extra trailing column means the xlsx schema has evolved and the
    loader needs a code update before the new data is trustworthy."""
    extra_header = (*_HEADER, "DefaultValue")
    path = _write_sheet(tmp_path / "tf.xlsx", [
        extra_header,
        (None, None, None, "T", "a", None, "Long", True, None, None, "newdata"),
    ])
    with pytest.raises(AccessSchemaError, match=r"unknown extra header"):
        load_access_schema(path)


def test_blank_header_cell_with_data_in_extra_column_raises(tmp_path: Path) -> None:
    """If the header's extra column is BLANK but a data row has a value
    past column 10, the loader must still raise — otherwise the data is
    silently dropped, violating the strict-schema contract."""
    header_with_blank_extra = (*_HEADER, None)
    path = _write_sheet(tmp_path / "tf.xlsx", [
        header_with_blank_extra,
        (None, None, None, "T", "a", None, "Long", True, None, None, "leaked"),
    ])
    with pytest.raises(AccessSchemaError, match=r"unexpected non-blank data"):
        load_access_schema(path)


def test_dump_side_names_preserved_when_present(tmp_path: Path) -> None:
    """A row that explicitly maps a Datadump table/field to a different
    Access name must surface both names — the mdb writer (1.3b) and the
    parity comparator (1.5) need the Datadump side to find source rows."""
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, "BIOG_MAIN_OLD", "c_personid_old", "BIOG_MAIN", "c_personid",
         "Primary Key", "Long", False, None, None),
    ])
    schema = load_access_schema(path)
    col = schema["BIOG_MAIN"].columns[0]
    assert col.dump_table_name == "BIOG_MAIN_OLD"
    assert col.dump_field_name == "c_personid_old"


def test_dump_side_names_default_none_when_xlsx_blank(tmp_path: Path) -> None:
    """The real xlsx has blank DumpTblNm/DumpFldNm in every row (implicit
    same-name mapping). Loader must store None, not the empty string."""
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "T", "a", None, "Long", True, None, None),
    ])
    schema = load_access_schema(path)
    col = schema["T"].columns[0]
    assert col.dump_table_name is None
    assert col.dump_field_name is None


def test_whitespace_only_table_or_column_is_skipped(tmp_path: Path) -> None:
    """A row where AccessTblNm/AccessFldNm contains only whitespace must
    NOT become an empty-named table in the schema."""
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "   ", "   ", None, "Long", True, None, None),
        (None, None, None, "T", "   ", None, "Long", True, None, None),
        (None, None, None, "T", "a", None, "Long", True, None, None),
    ])
    schema = load_access_schema(path)
    assert list(schema.tables.keys()) == ["T"]
    assert [c.name for c in schema["T"].columns] == ["a"]


def test_iterable_and_indexable(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / "tf.xlsx", [
        _HEADER,
        (None, None, None, "X", "a", None, "Long", True, None, None),
        (None, None, None, "Y", "b", None, "Long", True, None, None),
    ])
    schema = load_access_schema(path)
    names_via_iter = sorted(t.name for t in schema)
    assert names_via_iter == ["X", "Y"]
    assert schema["X"].name == "X"
